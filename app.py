from __future__ import annotations

import hashlib
import hmac
import json
import os
import secrets
import sqlite3
from contextlib import contextmanager
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from starlette.middleware.sessions import SessionMiddleware

BASE_DIR = Path(__file__).resolve().parent
PUBLIC_DIR = BASE_DIR / "public"
DB_PATH = BASE_DIR / "data.sqlite"
SECRET_KEY = os.environ.get("SESSION_SECRET", "teklif-takip-fastapi-secret")
PORT = int(os.environ.get("PORT", "3030"))

app = FastAPI(title="TeklifTakip")
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, max_age=60 * 60 * 8, same_site="lax")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])
app.mount("/public", StaticFiles(directory=PUBLIC_DIR), name="public")

OFFER_TYPES = ["vehicle", "transfer", "tour"]
OFFER_STATUSES = ["new", "quoted", "waiting", "returned", "won", "no_reply", "lost", "cancelled", "archived"]
CHANNELS = ["whatsapp", "mail", "office", "phone", "instagram", "website", "reference", "other"]
CURRENCIES = ["TL", "USD", "EUR", "GBP"]
PRICE_TYPES = ["daily", "total"]
TRIP_TYPES = ["one_way", "round_trip"]
ROLES = ["admin", "user"]

OFFER_TYPE_LABELS = {
    "vehicle": "Araç",
    "transfer": "Transfer",
    "tour": "Tur",
}
STATUS_LABELS = {
    "new": "Yeni Talep",
    "quoted": "Teklif Verildi",
    "waiting": "Beklemede",
    "returned": "Geri Döndü",
    "won": "Satış Oldu",
    "no_reply": "Cevap Yok",
    "lost": "Kaçtı",
    "cancelled": "İptal",
    "archived": "Arşiv",
}
CHANNEL_LABELS = {
    "whatsapp": "WhatsApp",
    "mail": "Mail",
    "office": "Ofis",
    "phone": "Telefon",
    "instagram": "Instagram",
    "website": "Web Site",
    "reference": "Referans",
    "other": "Diğer",
}
PRICE_TYPE_LABELS = {"daily": "Günlük", "total": "Toplam"}
TRIP_TYPE_LABELS = {"one_way": "Tek Yön", "round_trip": "Çift Yön"}
ROLE_LABELS = {"admin": "Admin", "user": "Kullanıcı"}

VEHICLE_SEED = [
    {"name": "Fiat Egea - Petrol | Manuel", "scope": "vehicle"},
    {"name": "Fiat Egea - Diesel | Automatic", "scope": "vehicle"},
    {"name": "Fiat Egea Cross - Petrol | Manuel", "scope": "vehicle"},
    {"name": "Renault Clio 5 HB - Petrol | Automatic", "scope": "vehicle"},
    {"name": "Opel Corsa - Petrol | Automatic", "scope": "vehicle"},
    {"name": "Opel Mokka - Petrol | Automatic", "scope": "vehicle"},
    {"name": "Peugeot 2008 - Petrol | Automatic", "scope": "vehicle"},
    {"name": "Dacia Lodgy (7 seats) - Diesel | Manuel", "scope": "vehicle"},
    {"name": "Mercedes Vito", "scope": "transfer"},
    {"name": "Mercedes Sprinter", "scope": "transfer"},
    {"name": "Dış Araç - Minibüs - Diesel | Manuel veya Automatic", "scope": "transfer"},
]


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def normalize_phone(phone: str) -> str:
    return "".join(str(phone or "").split()).strip()


def compute_30_day_rental_days(start_iso: str, end_iso: str) -> int:
    if not start_iso or not end_iso:
        return 0
    start = date.fromisoformat(start_iso)
    end = date.fromisoformat(end_iso)
    if end < start:
        return 0
    start_day = min(start.day, 30)
    end_day = min(end.day, 30)
    return ((end.year - start.year) * 360) + ((end.month - start.month) * 30) + (end_day - start_day) + 1


@contextmanager
def get_db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def row_to_dict(row: sqlite3.Row | None) -> dict[str, Any] | None:
    return dict(row) if row is not None else None


def rows_to_dicts(rows: list[sqlite3.Row]) -> list[dict[str, Any]]:
    return [dict(r) for r in rows]


def hash_password(password: str) -> str:
    salt = secrets.token_hex(16)
    key = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 120000).hex()
    return f"{salt}:{key}"


def verify_password(password: str, hashed: str) -> bool:
    try:
        salt, key = hashed.split(":", 1)
    except ValueError:
        return False
    candidate = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 120000).hex()
    return hmac.compare_digest(candidate, key)


def validate_choice(value: str, allowed: list[str], field: str):
    if value not in allowed:
        raise HTTPException(status_code=400, detail=f"{field} geçersiz.")


def has_column(conn: sqlite3.Connection, table: str, column: str) -> bool:
    rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
    return any(r[1] == column for r in rows)


def ensure_column(conn: sqlite3.Connection, table: str, column_name: str, definition: str):
    if not has_column(conn, table, column_name):
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {definition}")


def seed_vehicles(conn: sqlite3.Connection):
    ts = now_iso()
    for item in VEHICLE_SEED:
        row = conn.execute("SELECT id FROM vehicles WHERE name = ?", (item["name"],)).fetchone()
        if row:
            conn.execute(
                "UPDATE vehicles SET is_active = 1, scope = ?, updated_at = ? WHERE id = ?",
                (item["scope"], ts, row["id"]),
            )
        else:
            conn.execute(
                "INSERT INTO vehicles (name, scope, is_active, created_at, updated_at) VALUES (?, ?, 1, ?, ?)",
                (item["name"], item["scope"], ts, ts),
            )
    allowed_names = [v["name"] for v in VEHICLE_SEED]
    if allowed_names:
        placeholders = ",".join("?" for _ in allowed_names)
        conn.execute(
            f"UPDATE vehicles SET is_active = 0, updated_at = ? WHERE name NOT IN ({placeholders})",
            [ts, *allowed_names],
        )


def init_db():
    with get_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              username TEXT NOT NULL UNIQUE,
              display_name TEXT NOT NULL,
              password_hash TEXT NOT NULL,
              role TEXT NOT NULL,
              is_active INTEGER NOT NULL DEFAULT 1,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS vehicles (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              name TEXT NOT NULL UNIQUE,
              scope TEXT NOT NULL DEFAULT 'vehicle',
              is_active INTEGER NOT NULL DEFAULT 1,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS customers (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              full_name TEXT NOT NULL,
              phone TEXT NOT NULL UNIQUE,
              note TEXT DEFAULT '',
              is_active INTEGER NOT NULL DEFAULT 1,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL,
              last_contact_at TEXT
            );
            CREATE TABLE IF NOT EXISTS offers (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              customer_id INTEGER NOT NULL,
              created_by_user_id INTEGER NOT NULL,
              offer_type TEXT NOT NULL,
              request_channel TEXT NOT NULL,
              response_channel TEXT NOT NULL,
              currency TEXT NOT NULL,
              price_type TEXT NOT NULL,
              price_amount REAL NOT NULL,
              status TEXT NOT NULL,
              offer_date TEXT NOT NULL,
              note TEXT DEFAULT '',
              followup_date TEXT,
              is_archived INTEGER NOT NULL DEFAULT 0,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL,
              FOREIGN KEY (customer_id) REFERENCES customers(id) ON DELETE RESTRICT,
              FOREIGN KEY (created_by_user_id) REFERENCES users(id) ON DELETE RESTRICT
            );
            CREATE TABLE IF NOT EXISTS offer_vehicle_details (
              offer_id INTEGER PRIMARY KEY,
              vehicle_id INTEGER NOT NULL,
              transmission TEXT NOT NULL DEFAULT '',
              pickup_date TEXT NOT NULL,
              return_date TEXT NOT NULL,
              total_days INTEGER NOT NULL,
              pickup_location TEXT DEFAULT '',
              return_location TEXT DEFAULT '',
              FOREIGN KEY (offer_id) REFERENCES offers(id) ON DELETE CASCADE,
              FOREIGN KEY (vehicle_id) REFERENCES vehicles(id) ON DELETE RESTRICT
            );
            CREATE TABLE IF NOT EXISTS offer_transfer_details (
              offer_id INTEGER PRIMARY KEY,
              vehicle_id INTEGER,
              trip_type TEXT NOT NULL,
              pickup_location TEXT NOT NULL,
              dropoff_location TEXT NOT NULL,
              transfer_date TEXT NOT NULL,
              FOREIGN KEY (offer_id) REFERENCES offers(id) ON DELETE CASCADE,
              FOREIGN KEY (vehicle_id) REFERENCES vehicles(id) ON DELETE RESTRICT
            );
            CREATE TABLE IF NOT EXISTS offer_tour_details (
              offer_id INTEGER PRIMARY KEY,
              tour_name TEXT NOT NULL,
              participant_count INTEGER NOT NULL,
              tour_date TEXT NOT NULL,
              FOREIGN KEY (offer_id) REFERENCES offers(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS followups (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              offer_id INTEGER NOT NULL,
              followup_date TEXT NOT NULL,
              followup_status TEXT NOT NULL,
              note TEXT DEFAULT '',
              created_by_user_id INTEGER NOT NULL,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL,
              FOREIGN KEY (offer_id) REFERENCES offers(id) ON DELETE CASCADE,
              FOREIGN KEY (created_by_user_id) REFERENCES users(id) ON DELETE RESTRICT
            );
            CREATE TABLE IF NOT EXISTS audit_logs (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              user_id INTEGER,
              entity_type TEXT NOT NULL,
              entity_id TEXT NOT NULL,
              action TEXT NOT NULL,
              old_value_json TEXT,
              new_value_json TEXT,
              created_at TEXT NOT NULL,
              FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE SET NULL
            );
            """
        )
        ensure_column(conn, "vehicles", "scope", "scope TEXT NOT NULL DEFAULT 'vehicle'")
        ensure_column(conn, "offer_transfer_details", "vehicle_id", "vehicle_id INTEGER")

        user_count = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        if not user_count:
            ts = now_iso()
            conn.execute(
                "INSERT INTO users (username, display_name, password_hash, role, is_active, created_at, updated_at) VALUES (?, ?, ?, 'admin', 1, ?, ?)",
                ("admin", "Admin", hash_password("admin123"), ts, ts),
            )

        seed_vehicles(conn)


@app.on_event("startup")
def startup_event():
    init_db()


def current_user(request: Request) -> dict[str, Any]:
    user = request.session.get("user")
    if not user:
        raise HTTPException(status_code=401, detail="Oturum gerekli.")
    return user


def require_admin(request: Request) -> dict[str, Any]:
    user = current_user(request)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin yetkisi gerekli.")
    return user


def log_action(user_id: int | None, entity_type: str, entity_id: str | int, action: str, old_value: Any = None, new_value: Any = None):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO audit_logs (user_id, entity_type, entity_id, action, old_value_json, new_value_json, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                user_id,
                entity_type,
                str(entity_id),
                action,
                json.dumps(old_value, ensure_ascii=False) if old_value is not None else None,
                json.dumps(new_value, ensure_ascii=False) if new_value is not None else None,
                now_iso(),
            ),
        )


def upsert_customer(conn: sqlite3.Connection, full_name: str, phone: str, note: str, current_customer_id: int | None = None) -> int:
    full_name = str(full_name or "").strip()
    phone = normalize_phone(phone)
    note = str(note or "").strip()
    if not full_name:
        raise HTTPException(status_code=400, detail="Müşteri adı gerekli.")
    if not phone:
        raise HTTPException(status_code=400, detail="Telefon gerekli.")
    ts = now_iso()
    existing = conn.execute("SELECT * FROM customers WHERE phone = ?", (phone,)).fetchone()

    if current_customer_id:
        current_row = conn.execute("SELECT * FROM customers WHERE id = ?", (current_customer_id,)).fetchone()
        if not current_row:
            raise HTTPException(status_code=404, detail="Müşteri bulunamadı.")
        if existing and existing["id"] != current_customer_id:
            conn.execute(
                "UPDATE customers SET full_name = ?, note = ?, last_contact_at = ?, updated_at = ? WHERE id = ?",
                (full_name, note, ts, ts, existing["id"]),
            )
            return existing["id"]
        conn.execute(
            "UPDATE customers SET full_name = ?, phone = ?, note = ?, last_contact_at = ?, updated_at = ? WHERE id = ?",
            (full_name, phone, note, ts, ts, current_customer_id),
        )
        return current_customer_id

    if existing:
        conn.execute(
            "UPDATE customers SET full_name = ?, note = ?, last_contact_at = ?, updated_at = ? WHERE id = ?",
            (full_name, note, ts, ts, existing["id"]),
        )
        return existing["id"]

    cur = conn.execute(
        "INSERT INTO customers (full_name, phone, note, is_active, created_at, updated_at, last_contact_at) VALUES (?, ?, ?, 1, ?, ?, ?)",
        (full_name, phone, note, ts, ts, ts),
    )
    return cur.lastrowid


def replace_offer_details(conn: sqlite3.Connection, offer_id: int, offer_type: str, payload: dict[str, Any]):
    conn.execute("DELETE FROM offer_vehicle_details WHERE offer_id = ?", (offer_id,))
    conn.execute("DELETE FROM offer_transfer_details WHERE offer_id = ?", (offer_id,))
    conn.execute("DELETE FROM offer_tour_details WHERE offer_id = ?", (offer_id,))

    if offer_type == "vehicle":
        vehicle_id = int(payload.get("vehicle_id") or 0)
        pickup_date = str(payload.get("pickup_date", "")).strip()
        return_date = str(payload.get("return_date", "")).strip()
        pickup_location = str(payload.get("pickup_location", "")).strip()
        return_location = str(payload.get("return_location", "")).strip()
        total_days = int(payload.get("total_days") or 0)
        if not total_days and pickup_date and return_date:
            total_days = compute_30_day_rental_days(pickup_date, return_date)
        if not vehicle_id or not pickup_date or not return_date or total_days <= 0:
            raise HTTPException(status_code=400, detail="Araç detayları eksik veya hatalı.")
        conn.execute(
            "INSERT INTO offer_vehicle_details (offer_id, vehicle_id, transmission, pickup_date, return_date, total_days, pickup_location, return_location) VALUES (?, ?, '', ?, ?, ?, ?, ?)",
            (offer_id, vehicle_id, pickup_date, return_date, total_days, pickup_location, return_location),
        )

    elif offer_type == "transfer":
        trip_type = str(payload.get("trip_type", "")).strip()
        transfer_date = str(payload.get("transfer_date", "")).strip()
        pickup_location = str(payload.get("transfer_pickup_location", "")).strip()
        dropoff_location = str(payload.get("transfer_dropoff_location", "")).strip()
        vehicle_id = int(payload.get("transfer_vehicle_id") or 0) or None
        validate_choice(trip_type, TRIP_TYPES, "Transfer tipi")
        if not transfer_date or not pickup_location or not dropoff_location:
            raise HTTPException(status_code=400, detail="Transfer detayları eksik.")
        conn.execute(
            "INSERT INTO offer_transfer_details (offer_id, vehicle_id, trip_type, pickup_location, dropoff_location, transfer_date) VALUES (?, ?, ?, ?, ?, ?)",
            (offer_id, vehicle_id, trip_type, pickup_location, dropoff_location, transfer_date),
        )

    elif offer_type == "tour":
        tour_name = str(payload.get("tour_name", "")).strip()
        participant_count = int(payload.get("participant_count") or 0)
        tour_date = str(payload.get("tour_date", "")).strip()
        if not tour_name or participant_count <= 0 or not tour_date:
            raise HTTPException(status_code=400, detail="Tur detayları eksik.")
        conn.execute(
            "INSERT INTO offer_tour_details (offer_id, tour_name, participant_count, tour_date) VALUES (?, ?, ?, ?)",
            (offer_id, tour_name, participant_count, tour_date),
        )


def get_offer_detail(offer_id: int) -> dict[str, Any] | None:
    with get_db() as conn:
        offer = conn.execute(
            """
            SELECT o.*, c.full_name AS customer_name, c.phone AS customer_phone, c.note AS customer_note,
                   u.display_name AS created_by_name
            FROM offers o
            JOIN customers c ON c.id = o.customer_id
            JOIN users u ON u.id = o.created_by_user_id
            WHERE o.id = ?
            """,
            (offer_id,),
        ).fetchone()
        if not offer:
            return None
        data = row_to_dict(offer)
        detail = None
        if data["offer_type"] == "vehicle":
            detail = row_to_dict(conn.execute(
                """
                SELECT d.*, v.name AS vehicle_name, v.scope AS vehicle_scope
                FROM offer_vehicle_details d
                JOIN vehicles v ON v.id = d.vehicle_id
                WHERE d.offer_id = ?
                """,
                (offer_id,),
            ).fetchone())
        elif data["offer_type"] == "transfer":
            detail = row_to_dict(conn.execute(
                """
                SELECT d.*, v.name AS vehicle_name
                FROM offer_transfer_details d
                LEFT JOIN vehicles v ON v.id = d.vehicle_id
                WHERE d.offer_id = ?
                """,
                (offer_id,),
            ).fetchone())
        elif data["offer_type"] == "tour":
            detail = row_to_dict(conn.execute("SELECT * FROM offer_tour_details WHERE offer_id = ?", (offer_id,)).fetchone())
        data["detail"] = detail
        return data


@app.get("/")
def root():
    return FileResponse(PUBLIC_DIR / "index.html")


@app.get("/sw.js")
def service_worker():
    return FileResponse(PUBLIC_DIR / "sw.js")


@app.get("/manifest.webmanifest")
def manifest():
    return FileResponse(PUBLIC_DIR / "manifest.webmanifest")


@app.post("/api/login")
async def login(request: Request):
    payload = await request.json()
    username = str(payload.get("username", "")).strip()
    password = str(payload.get("password", ""))
    with get_db() as conn:
        row = conn.execute("SELECT * FROM users WHERE username = ? AND is_active = 1", (username,)).fetchone()
        if not row or not verify_password(password, row["password_hash"]):
            raise HTTPException(status_code=401, detail="Kullanıcı adı veya şifre hatalı.")
        user = {"id": row["id"], "username": row["username"], "display_name": row["display_name"], "role": row["role"]}
        request.session["user"] = user
        log_action(row["id"], "session", row["id"], "login", None, {"username": row["username"]})
        return {"user": user}


@app.post("/api/logout")
def logout(request: Request):
    user = current_user(request)
    request.session.clear()
    log_action(user["id"], "session", user["id"], "logout", None, {"username": user["username"]})
    return {"ok": True}


@app.get("/api/session")
def session_info(request: Request):
    return {"user": request.session.get("user")}


@app.get("/api/meta")
def meta(request: Request):
    current_user(request)
    return {
        "offerTypes": OFFER_TYPES,
        "offerStatuses": OFFER_STATUSES,
        "channels": CHANNELS,
        "currencies": CURRENCIES,
        "priceTypes": PRICE_TYPES,
        "tripTypes": TRIP_TYPES,
        "roles": ROLES,
        "labels": {
            "offerTypes": OFFER_TYPE_LABELS,
            "offerStatuses": STATUS_LABELS,
            "channels": CHANNEL_LABELS,
            "priceTypes": PRICE_TYPE_LABELS,
            "tripTypes": TRIP_TYPE_LABELS,
            "roles": ROLE_LABELS,
        },
    }


@app.get("/api/dashboard")
def dashboard(request: Request):
    current_user(request)
    today = date.today().isoformat()
    with get_db() as conn:
        today_offers = conn.execute("SELECT COUNT(*) FROM offers WHERE offer_date = ?", (today,)).fetchone()[0]
        waiting_offers = conn.execute("SELECT COUNT(*) FROM offers WHERE status IN ('new','quoted','waiting','returned') AND is_archived = 0").fetchone()[0]
        won_offers = conn.execute("SELECT COUNT(*) FROM offers WHERE status = 'won' AND is_archived = 0").fetchone()[0]
        lost_offers = conn.execute("SELECT COUNT(*) FROM offers WHERE status = 'lost' AND is_archived = 0").fetchone()[0]
        latest_offers = rows_to_dicts(conn.execute(
            """
            SELECT o.id, o.offer_type, o.status, o.offer_date, o.price_amount, o.currency,
                   c.full_name AS customer_name, u.display_name AS created_by_name
            FROM offers o
            JOIN customers c ON c.id = o.customer_id
            JOIN users u ON u.id = o.created_by_user_id
            ORDER BY o.created_at DESC LIMIT 8
            """
        ).fetchall())
        latest_customers = rows_to_dicts(conn.execute(
            "SELECT id, full_name, phone, created_at FROM customers ORDER BY created_at DESC LIMIT 8"
        ).fetchall())
    return {
        "todayOffers": today_offers,
        "waitingOffers": waiting_offers,
        "wonOffers": won_offers,
        "lostOffers": lost_offers,
        "latestOffers": latest_offers,
        "latestCustomers": latest_customers,
    }


@app.get("/api/customers")
def customers(request: Request, search: str = ""):
    current_user(request)
    term = f"%{search}%"
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT c.*, COUNT(o.id) AS offers_count, MAX(o.offer_date) AS latest_offer_date
            FROM customers c
            LEFT JOIN offers o ON o.customer_id = c.id
            WHERE c.full_name LIKE ? OR c.phone LIKE ?
            GROUP BY c.id
            ORDER BY c.updated_at DESC
            LIMIT 300
            """,
            (term, term),
        ).fetchall()
    return rows_to_dicts(rows)


@app.post("/api/customers")
async def create_customer(request: Request):
    user = current_user(request)
    payload = await request.json()
    ts = now_iso()
    with get_db() as conn:
        customer_id = upsert_customer(conn, payload.get("full_name", ""), payload.get("phone", ""), payload.get("note", ""), None)
        row = conn.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchone()
    item = row_to_dict(row)
    log_action(user["id"], "customer", customer_id, "create_or_update", None, item)
    return item


@app.put("/api/customers/{customer_id}")
async def update_customer(customer_id: int, request: Request):
    user = current_user(request)
    payload = await request.json()
    with get_db() as conn:
        before = conn.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchone()
        if not before:
            raise HTTPException(status_code=404, detail="Müşteri bulunamadı.")
        upsert_customer(conn, payload.get("full_name", ""), payload.get("phone", ""), payload.get("note", ""), customer_id)
        is_active = 1 if payload.get("is_active") in (1, True, "1", "true", "on") else 0
        conn.execute("UPDATE customers SET is_active = ?, updated_at = ? WHERE id = ?", (is_active, now_iso(), customer_id))
        after = conn.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchone()
    log_action(user["id"], "customer", customer_id, "update", row_to_dict(before), row_to_dict(after))
    return row_to_dict(after)


@app.get("/api/customers/{customer_id}/offers")
def customer_offers(customer_id: int, request: Request):
    current_user(request)
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT o.id, o.offer_type, o.status, o.price_type, o.price_amount, o.currency, o.offer_date,
                   u.display_name AS created_by_name
            FROM offers o
            JOIN users u ON u.id = o.created_by_user_id
            WHERE o.customer_id = ?
            ORDER BY o.created_at DESC
            """,
            (customer_id,),
        ).fetchall()
    return rows_to_dicts(rows)


@app.get("/api/offers")
def offers(
    request: Request,
    search: str = "",
    status: str = "",
    offer_type: str = "",
    created_by_user_id: str = "",
    vehicle_id: str = "",
    from_date: str = "",
    to: str = "",
    from_: str | None = None,
):
    current_user(request)
    from_value = from_ or from_date or request.query_params.get("from", "")
    clauses = ["1 = 1"]
    params: list[Any] = []
    if search:
        clauses.append("(c.full_name LIKE ? OR c.phone LIKE ? OR IFNULL(vv.name, '') LIKE ? OR IFNULL(tv.name, '') LIKE ? OR IFNULL(t.tour_name, '') LIKE ? OR IFNULL(tr.pickup_location, '') LIKE ? OR IFNULL(tr.dropoff_location, '') LIKE ?)")
        params.extend([f"%{search}%"] * 7)
    if status:
        clauses.append("o.status = ?")
        params.append(status)
    if offer_type:
        clauses.append("o.offer_type = ?")
        params.append(offer_type)
    if created_by_user_id:
        clauses.append("o.created_by_user_id = ?")
        params.append(int(created_by_user_id))
    if vehicle_id:
        clauses.append("(vd.vehicle_id = ? OR tr.vehicle_id = ?)")
        params.extend([int(vehicle_id), int(vehicle_id)])
    if from_value:
        clauses.append("o.offer_date >= ?")
        params.append(from_value)
    if to:
        clauses.append("o.offer_date <= ?")
        params.append(to)
    with get_db() as conn:
        rows = conn.execute(
            f"""
            SELECT o.id, o.offer_type, o.status, o.offer_date, o.request_channel, o.price_type, o.price_amount, o.currency,
                   c.full_name AS customer_name, c.phone AS customer_phone,
                   u.display_name AS created_by_name,
                   vv.name AS vehicle_name,
                   vd.pickup_date, vd.return_date, vd.total_days, vd.pickup_location, vd.return_location,
                   tv.name AS transfer_vehicle_name,
                   tr.trip_type, tr.pickup_location AS transfer_pickup, tr.dropoff_location AS transfer_dropoff, tr.transfer_date,
                   t.tour_name, t.participant_count, t.tour_date
            FROM offers o
            JOIN customers c ON c.id = o.customer_id
            JOIN users u ON u.id = o.created_by_user_id
            LEFT JOIN offer_vehicle_details vd ON vd.offer_id = o.id
            LEFT JOIN vehicles vv ON vv.id = vd.vehicle_id
            LEFT JOIN offer_transfer_details tr ON tr.offer_id = o.id
            LEFT JOIN vehicles tv ON tv.id = tr.vehicle_id
            LEFT JOIN offer_tour_details t ON t.offer_id = o.id
            WHERE {' AND '.join(clauses)}
            ORDER BY o.created_at DESC
            LIMIT 500
            """,
            params,
        ).fetchall()
    return rows_to_dicts(rows)


@app.post("/api/offers")
async def create_offer(request: Request):
    user = current_user(request)
    payload = await request.json()
    offer_type = str(payload.get("offer_type", "")).strip()
    channel = str(payload.get("channel", payload.get("request_channel", ""))).strip()
    currency = str(payload.get("currency", "")).strip()
    price_type = str(payload.get("price_type", "")).strip()
    status = str(payload.get("status", "")).strip()
    offer_date = str(payload.get("offer_date", "")).strip()
    note = str(payload.get("note", "")).strip()
    customer_name = str(payload.get("customer_full_name", "")).strip()
    customer_phone = normalize_phone(payload.get("customer_phone", ""))
    customer_note = str(payload.get("customer_note", "")).strip()
    try:
        price_amount = float(payload.get("price_amount") or 0)
    except ValueError:
        price_amount = 0
    if not offer_date or price_amount <= 0:
        raise HTTPException(status_code=400, detail="Teklif tarihi ve fiyat gerekli.")
    validate_choice(offer_type, OFFER_TYPES, "Teklif tipi")
    validate_choice(channel, CHANNELS, "Kanal")
    validate_choice(currency, CURRENCIES, "Para birimi")
    validate_choice(price_type, PRICE_TYPES, "Fiyat tipi")
    validate_choice(status, OFFER_STATUSES, "Durum")

    ts = now_iso()
    with get_db() as conn:
        customer_id = upsert_customer(conn, customer_name, customer_phone, customer_note)
        cur = conn.execute(
            """
            INSERT INTO offers (customer_id, created_by_user_id, offer_type, request_channel, response_channel, currency, price_type, price_amount, status, offer_date, note, followup_date, is_archived, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, ?, ?, ?)
            """,
            (customer_id, user["id"], offer_type, channel, channel, currency, price_type, price_amount, status, offer_date, note, 1 if status == "archived" else 0, ts, ts),
        )
        offer_id = cur.lastrowid
        replace_offer_details(conn, offer_id, offer_type, payload)
    detail = get_offer_detail(offer_id)
    log_action(user["id"], "offer", offer_id, "create", None, detail)
    return detail


@app.get("/api/offers/{offer_id}")
def offer_detail(offer_id: int, request: Request):
    current_user(request)
    detail = get_offer_detail(offer_id)
    if not detail:
        raise HTTPException(status_code=404, detail="Teklif bulunamadı.")
    return detail


@app.delete("/api/offers/{offer_id}")
def delete_offer(offer_id: int, request: Request):
    user = current_user(request)
    before = get_offer_detail(offer_id)
    if not before:
        raise HTTPException(status_code=404, detail="Teklif bulunamadı.")
    with get_db() as conn:
        cur = conn.execute("DELETE FROM offers WHERE id = ?", (offer_id,))
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="Teklif bulunamadı.")
    log_action(user["id"], "offer", offer_id, "delete", before, None)
    return {"ok": True, "deleted_id": offer_id}


@app.put("/api/offers/{offer_id}")
async def update_offer(offer_id: int, request: Request):
    user = current_user(request)
    payload = await request.json()
    before = get_offer_detail(offer_id)
    if not before:
        raise HTTPException(status_code=404, detail="Teklif bulunamadı.")
    channel = str(payload.get("channel", payload.get("request_channel", ""))).strip()
    currency = str(payload.get("currency", "")).strip()
    price_type = str(payload.get("price_type", "")).strip()
    status = str(payload.get("status", "")).strip()
    offer_date = str(payload.get("offer_date", "")).strip()
    note = str(payload.get("note", "")).strip()
    customer_name = str(payload.get("customer_full_name", before.get("customer_name", ""))).strip()
    customer_phone = normalize_phone(payload.get("customer_phone", before.get("customer_phone", "")))
    customer_note = str(payload.get("customer_note", before.get("customer_note", ""))).strip()
    try:
        price_amount = float(payload.get("price_amount") or 0)
    except ValueError:
        price_amount = 0
    if not offer_date or price_amount <= 0:
        raise HTTPException(status_code=400, detail="Teklif tarihi ve fiyat gerekli.")
    validate_choice(channel, CHANNELS, "Kanal")
    validate_choice(currency, CURRENCIES, "Para birimi")
    validate_choice(price_type, PRICE_TYPES, "Fiyat tipi")
    validate_choice(status, OFFER_STATUSES, "Durum")

    with get_db() as conn:
        customer_id = upsert_customer(conn, customer_name, customer_phone, customer_note, before["customer_id"])
        conn.execute(
            """
            UPDATE offers SET customer_id = ?, request_channel = ?, response_channel = ?, currency = ?, price_type = ?, price_amount = ?, status = ?, offer_date = ?, note = ?, updated_at = ?, is_archived = ?
            WHERE id = ?
            """,
            (customer_id, channel, channel, currency, price_type, price_amount, status, offer_date, note, now_iso(), 1 if status == "archived" else 0, offer_id),
        )
        replace_offer_details(conn, offer_id, before["offer_type"], payload)
    after = get_offer_detail(offer_id)
    log_action(user["id"], "offer", offer_id, "update", before, after)
    return after


@app.get("/api/users")
def users(request: Request):
    current_user(request)
    with get_db() as conn:
        rows = conn.execute("SELECT id, username, display_name, role, is_active, created_at, updated_at FROM users ORDER BY display_name ASC").fetchall()
    return rows_to_dicts(rows)


@app.post("/api/users")
async def create_user(request: Request):
    admin = require_admin(request)
    payload = await request.json()
    username = str(payload.get("username", "")).strip()
    display_name = str(payload.get("display_name", "")).strip()
    role = str(payload.get("role", "user")).strip()
    password = str(payload.get("password", "")).strip()
    is_active = 1 if payload.get("is_active") in (1, True, "1", "true", "on") else 0
    if not username or not display_name or not password:
        raise HTTPException(status_code=400, detail="Kullanıcı adı, görünen ad ve şifre gerekli.")
    validate_choice(role, ROLES, "Rol")
    ts = now_iso()
    with get_db() as conn:
        try:
            cur = conn.execute(
                "INSERT INTO users (username, display_name, password_hash, role, is_active, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (username, display_name, hash_password(password), role, is_active, ts, ts),
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=400, detail="Bu kullanıcı adı zaten var.")
        row = conn.execute("SELECT id, username, display_name, role, is_active, created_at, updated_at FROM users WHERE id = ?", (cur.lastrowid,)).fetchone()
    item = row_to_dict(row)
    log_action(admin["id"], "user", item["id"], "create", None, item)
    return item


@app.put("/api/users/{user_id}")
async def update_user(user_id: int, request: Request):
    admin = require_admin(request)
    payload = await request.json()
    display_name = str(payload.get("display_name", "")).strip()
    role = str(payload.get("role", "user")).strip()
    is_active = 1 if payload.get("is_active") in (1, True, "1", "true", "on") else 0
    if not display_name:
        raise HTTPException(status_code=400, detail="Görünen ad gerekli.")
    validate_choice(role, ROLES, "Rol")
    with get_db() as conn:
        before = conn.execute("SELECT id, username, display_name, role, is_active FROM users WHERE id = ?", (user_id,)).fetchone()
        if not before:
            raise HTTPException(status_code=404, detail="Kullanıcı bulunamadı.")
        conn.execute("UPDATE users SET display_name = ?, role = ?, is_active = ?, updated_at = ? WHERE id = ?", (display_name, role, is_active, now_iso(), user_id))
        after = conn.execute("SELECT id, username, display_name, role, is_active FROM users WHERE id = ?", (user_id,)).fetchone()
    log_action(admin["id"], "user", user_id, "update", row_to_dict(before), row_to_dict(after))
    return row_to_dict(after)


@app.put("/api/users/{user_id}/reset-password")
async def reset_password(user_id: int, request: Request):
    admin = require_admin(request)
    payload = await request.json()
    password = str(payload.get("password", "")).strip()
    if not password:
        raise HTTPException(status_code=400, detail="Şifre gerekli.")
    with get_db() as conn:
        before = conn.execute("SELECT id, username, display_name FROM users WHERE id = ?", (user_id,)).fetchone()
        if not before:
            raise HTTPException(status_code=404, detail="Kullanıcı bulunamadı.")
        conn.execute("UPDATE users SET password_hash = ?, updated_at = ? WHERE id = ?", (hash_password(password), now_iso(), user_id))
    log_action(admin["id"], "user", user_id, "reset_password", row_to_dict(before), {"reset_by": admin["username"]})
    return {"ok": True}


@app.get("/api/vehicles")
def vehicles(request: Request):
    current_user(request)
    scope = request.query_params.get("scope", "").strip()
    clauses = ["1 = 1"]
    params: list[Any] = []
    if scope:
        clauses.append("scope = ?")
        params.append(scope)
    with get_db() as conn:
        rows = conn.execute(f"SELECT * FROM vehicles WHERE {' AND '.join(clauses)} ORDER BY is_active DESC, id ASC", params).fetchall()
    return rows_to_dicts(rows)


@app.post("/api/vehicles")
async def create_vehicle(request: Request):
    admin = require_admin(request)
    payload = await request.json()
    name = str(payload.get("name", "")).strip()
    scope = str(payload.get("scope", "vehicle")).strip() or "vehicle"
    is_active = 1 if payload.get("is_active") in (1, True, "1", "true", "on") else 0
    if not name:
        raise HTTPException(status_code=400, detail="Araç adı gerekli.")
    if scope not in ("vehicle", "transfer"):
        raise HTTPException(status_code=400, detail="Araç tipi geçersiz.")
    ts = now_iso()
    with get_db() as conn:
        try:
            cur = conn.execute("INSERT INTO vehicles (name, scope, is_active, created_at, updated_at) VALUES (?, ?, ?, ?, ?)", (name, scope, is_active, ts, ts))
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=400, detail="Bu araç zaten listede var.")
        row = conn.execute("SELECT * FROM vehicles WHERE id = ?", (cur.lastrowid,)).fetchone()
    item = row_to_dict(row)
    log_action(admin["id"], "vehicle", item["id"], "create", None, item)
    return item


@app.put("/api/vehicles/{vehicle_id}")
async def update_vehicle(vehicle_id: int, request: Request):
    admin = require_admin(request)
    payload = await request.json()
    name = str(payload.get("name", "")).strip()
    scope = str(payload.get("scope", "vehicle")).strip() or "vehicle"
    is_active = 1 if payload.get("is_active") in (1, True, "1", "true", "on") else 0
    if not name:
        raise HTTPException(status_code=400, detail="Araç adı gerekli.")
    if scope not in ("vehicle", "transfer"):
        raise HTTPException(status_code=400, detail="Araç tipi geçersiz.")
    with get_db() as conn:
        before = conn.execute("SELECT * FROM vehicles WHERE id = ?", (vehicle_id,)).fetchone()
        if not before:
            raise HTTPException(status_code=404, detail="Araç bulunamadı.")
        try:
            conn.execute("UPDATE vehicles SET name = ?, scope = ?, is_active = ?, updated_at = ? WHERE id = ?", (name, scope, is_active, now_iso(), vehicle_id))
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=400, detail="Bu araç adı zaten kullanılıyor.")
        after = conn.execute("SELECT * FROM vehicles WHERE id = ?", (vehicle_id,)).fetchone()
    log_action(admin["id"], "vehicle", vehicle_id, "update", row_to_dict(before), row_to_dict(after))
    return row_to_dict(after)


@app.get("/api/reports/summary")
def reports_summary(request: Request):
    require_admin(request)
    with get_db() as conn:
        summary = {
            "totalOffers": conn.execute("SELECT COUNT(*) FROM offers").fetchone()[0],
            "wonOffers": conn.execute("SELECT COUNT(*) FROM offers WHERE status = 'won'").fetchone()[0],
            "lostOffers": conn.execute("SELECT COUNT(*) FROM offers WHERE status = 'lost'").fetchone()[0],
            "noReplyOffers": conn.execute("SELECT COUNT(*) FROM offers WHERE status = 'no_reply'").fetchone()[0],
            "waitingOffers": conn.execute("SELECT COUNT(*) FROM offers WHERE status IN ('new','quoted','waiting','returned')").fetchone()[0],
        }
        by_type = rows_to_dicts(conn.execute("SELECT offer_type, COUNT(*) AS count FROM offers GROUP BY offer_type ORDER BY count DESC").fetchall())
        by_user = rows_to_dicts(conn.execute("SELECT u.display_name, COUNT(o.id) AS count FROM users u LEFT JOIN offers o ON o.created_by_user_id = u.id GROUP BY u.id ORDER BY count DESC, u.display_name ASC").fetchall())
        by_channel = rows_to_dicts(conn.execute("SELECT request_channel, COUNT(*) AS count FROM offers GROUP BY request_channel ORDER BY count DESC").fetchall())
    return {"summary": summary, "byType": by_type, "byUser": by_user, "byChannel": by_channel}


def offers_for_export(filters: dict[str, str]) -> list[dict[str, Any]]:
    clauses = ["1 = 1"]
    params: list[Any] = []
    if filters.get("status"):
        clauses.append("o.status = ?")
        params.append(filters["status"])
    if filters.get("offer_type"):
        clauses.append("o.offer_type = ?")
        params.append(filters["offer_type"])
    if filters.get("from"):
        clauses.append("o.offer_date >= ?")
        params.append(filters["from"])
    if filters.get("to"):
        clauses.append("o.offer_date <= ?")
        params.append(filters["to"])
    with get_db() as conn:
        rows = conn.execute(
            f"""
            SELECT o.id AS TeklifNo,
                   c.full_name AS Musteri,
                   c.phone AS Telefon,
                   o.offer_type AS TeklifTuruKod,
                   o.request_channel AS KanalKod,
                   o.offer_date AS TeklifTarihi,
                   o.status AS DurumKod,
                   o.price_type AS FiyatTipiKod,
                   o.price_amount AS Fiyat,
                   o.currency AS ParaBirimi,
                   u.display_name AS TeklifiVeren,
                   vv.name AS Arac,
                   vd.pickup_date AS AlisTarihi,
                   vd.return_date AS DonusTarihi,
                   vd.total_days AS Gun,
                   vd.pickup_location AS TeslimYeri,
                   vd.return_location AS DonusYeri,
                   tv.name AS TransferAraci,
                   tr.trip_type AS TransferTipiKod,
                   tr.pickup_location AS TransferAlis,
                   tr.dropoff_location AS TransferVaris,
                   tr.transfer_date AS TransferTarihi,
                   t.tour_name AS TurAdi,
                   t.participant_count AS KisiSayisi,
                   t.tour_date AS TurTarihi,
                   o.note AS TeklifNotu
            FROM offers o
            JOIN customers c ON c.id = o.customer_id
            JOIN users u ON u.id = o.created_by_user_id
            LEFT JOIN offer_vehicle_details vd ON vd.offer_id = o.id
            LEFT JOIN vehicles vv ON vv.id = vd.vehicle_id
            LEFT JOIN offer_transfer_details tr ON tr.offer_id = o.id
            LEFT JOIN vehicles tv ON tv.id = tr.vehicle_id
            LEFT JOIN offer_tour_details t ON t.offer_id = o.id
            WHERE {' AND '.join(clauses)}
            ORDER BY o.created_at DESC
            """,
            params,
        ).fetchall()
    out = []
    for row in rows_to_dicts(rows):
        row["TeklifTuru"] = OFFER_TYPE_LABELS.get(row.pop("TeklifTuruKod"), "-")
        row["Kanal"] = CHANNEL_LABELS.get(row.pop("KanalKod"), "-")
        row["Durum"] = STATUS_LABELS.get(row.pop("DurumKod"), "-")
        row["FiyatTipi"] = PRICE_TYPE_LABELS.get(row.pop("FiyatTipiKod"), "-")
        row["TransferTipi"] = TRIP_TYPE_LABELS.get(row.pop("TransferTipiKod"), "-") if row.get("TransferTipiKod") else ""
        row.pop("TransferTipiKod", None)
        out.append(row)
    return out


def workbook_response(rows: list[dict[str, Any]], sheet_name: str, filename: str) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
    else:
        ws.append(["Veri Yok"])
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 40)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/reports/offers.xlsx")
def export_offers(request: Request, status: str = "", offer_type: str = "", from_date: str = "", to: str = "", from_: str | None = None):
    require_admin(request)
    rows = offers_for_export({"status": status, "offer_type": offer_type, "from": from_ or from_date or request.query_params.get("from", ""), "to": to})
    return workbook_response(rows, "Teklifler", "teklifler.xlsx")


@app.get("/api/reports/customers.xlsx")
def export_customers(request: Request):
    require_admin(request)
    with get_db() as conn:
        rows = rows_to_dicts(conn.execute(
            """
            SELECT c.id AS MusteriNo, c.full_name AS AdSoyad, c.phone AS Telefon, c.note AS Not,
                   COUNT(o.id) AS TeklifSayisi, MAX(o.offer_date) AS SonTeklif,
                   c.last_contact_at AS SonIletisim, CASE WHEN c.is_active = 1 THEN 'Evet' ELSE 'Hayır' END AS Aktif
            FROM customers c
            LEFT JOIN offers o ON o.customer_id = c.id
            GROUP BY c.id
            ORDER BY c.updated_at DESC
            """
        ).fetchall())
    return workbook_response(rows, "Musteriler", "musteriler.xlsx")


@app.get("/api/backup/export.json")
def export_backup(request: Request):
    require_admin(request)
    with get_db() as conn:
        backup = {
            "users": rows_to_dicts(conn.execute("SELECT * FROM users").fetchall()),
            "vehicles": rows_to_dicts(conn.execute("SELECT * FROM vehicles").fetchall()),
            "customers": rows_to_dicts(conn.execute("SELECT * FROM customers").fetchall()),
            "offers": rows_to_dicts(conn.execute("SELECT * FROM offers").fetchall()),
            "offerVehicleDetails": rows_to_dicts(conn.execute("SELECT * FROM offer_vehicle_details").fetchall()),
            "offerTransferDetails": rows_to_dicts(conn.execute("SELECT * FROM offer_transfer_details").fetchall()),
            "offerTourDetails": rows_to_dicts(conn.execute("SELECT * FROM offer_tour_details").fetchall()),
            "followups": rows_to_dicts(conn.execute("SELECT * FROM followups").fetchall()) if conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='followups'").fetchone() else [],
            "auditLogs": rows_to_dicts(conn.execute("SELECT * FROM audit_logs").fetchall()),
            "exportedAt": now_iso(),
        }
    return JSONResponse(content=backup, headers={"Content-Disposition": 'attachment; filename="teklif-takip-backup.json"'})


@app.post("/api/backup/import")
async def import_backup(request: Request, backup: UploadFile = File(...)):
    admin = require_admin(request)
    try:
        payload = json.loads((await backup.read()).decode("utf-8"))
    except Exception:
        raise HTTPException(status_code=400, detail="Yedek dosyası okunamadı.")

    with get_db() as conn:
        conn.execute("DELETE FROM offer_vehicle_details")
        conn.execute("DELETE FROM offer_transfer_details")
        conn.execute("DELETE FROM offer_tour_details")
        if conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='followups'").fetchone():
            conn.execute("DELETE FROM followups")
        conn.execute("DELETE FROM offers")
        conn.execute("DELETE FROM customers")
        conn.execute("DELETE FROM vehicles")
        conn.execute("DELETE FROM audit_logs")
        conn.execute("DELETE FROM users")

        def insert_many(table: str, rows: list[dict[str, Any]]):
            if not rows:
                return
            allowed_cols = [r[1] for r in conn.execute(f"PRAGMA table_info({table})").fetchall()]
            filtered = []
            for row in rows:
                item = {k: row[k] for k in row.keys() if k in allowed_cols}
                filtered.append(item)
            if not filtered:
                return
            keys = list(filtered[0].keys())
            placeholders = ",".join("?" for _ in keys)
            conn.executemany(
                f"INSERT INTO {table} ({','.join(keys)}) VALUES ({placeholders})",
                [[item.get(k) for k in keys] for item in filtered],
            )

        insert_many("users", payload.get("users", []))
        insert_many("vehicles", payload.get("vehicles", []))
        insert_many("customers", payload.get("customers", []))
        insert_many("offers", payload.get("offers", []))
        insert_many("offer_vehicle_details", payload.get("offerVehicleDetails", []))
        insert_many("offer_transfer_details", payload.get("offerTransferDetails", []))
        insert_many("offer_tour_details", payload.get("offerTourDetails", []))
        if conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='followups'").fetchone():
            insert_many("followups", payload.get("followups", []))
        insert_many("audit_logs", payload.get("auditLogs", []))

        if conn.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
            ts = now_iso()
            conn.execute(
                "INSERT INTO users (username, display_name, password_hash, role, is_active, created_at, updated_at) VALUES (?, ?, ?, 'admin', 1, ?, ?)",
                ("admin", "Admin", hash_password("admin123"), ts, ts),
            )
        seed_vehicles(conn)

    log_action(admin["id"], "backup", "system", "import", None, {"importedAt": now_iso()})
    return {"ok": True}


@app.get("/api/audit-logs")
def audit_logs(request: Request):
    require_admin(request)
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT a.*, u.display_name AS user_name
            FROM audit_logs a
            LEFT JOIN users u ON u.id = a.user_id
            ORDER BY a.created_at DESC
            LIMIT 300
            """
        ).fetchall()
    return rows_to_dicts(rows)


@app.exception_handler(HTTPException)
async def http_exception_handler(_: Request, exc: HTTPException):
    return JSONResponse(status_code=exc.status_code, content={"error": exc.detail})


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="0.0.0.0", port=PORT, reload=False)
